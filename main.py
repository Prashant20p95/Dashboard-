"""
OpsCommand — Main Application Entry Point
Run:  python main.py  →  http://127.0.0.1:8000
"""
import os
import config

from fastapi import FastAPI
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse

from database import init_db
from seed_data import seed_collateral_items

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# ── Create application ──────────────────────────────────────────────────────
app = FastAPI(
    title=config.PORTAL_TITLE,
    version=config.APP_VERSION,
    description=config.PORTAL_SUBTITLE,
)

@app.middleware("http")
async def add_no_cache_headers(request, call_next):
    response = await call_next(request)
    response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate, max-age=0"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    return response

# ── Create required directories ─────────────────────────────────────────────
for d in [
    "static",
    "uploads/imports",
    "uploads/exports",
    "uploads/invoices/printing",
    "uploads/invoices/courier",
    "uploads/templates",
]:
    os.makedirs(os.path.join(BASE_DIR, d), exist_ok=True)

# ── Include API routers ─────────────────────────────────────────────────────
from routes.events import router as events_router
from routes.participants import router as participants_router
from routes.organisers import router as organisers_router
from routes.printing import router as printing_router
from routes.stock import router as stock_router
from routes.dispatch import router as dispatch_router
from routes.tracking import router as tracking_router
from routes.billing import router as billing_router
from routes.import_engine import router as import_router
from routes.export_engine import router as export_router
from routes.dashboard import router as dashboard_router
from routes.audit import router as audit_router
from routes.datastore import router as datastore_router
from routes.config import router as config_router

app.include_router(config_router, prefix="/api/config", tags=["Configuration"])
app.include_router(events_router, prefix="/api/events", tags=["Events"])
app.include_router(participants_router, prefix="/api/participants", tags=["Participants"])
app.include_router(organisers_router, prefix="/api/organisers", tags=["Organisers"])
app.include_router(printing_router, prefix="/api/printing", tags=["Printing"])
app.include_router(stock_router, prefix="/api/stock", tags=["Stock"])
app.include_router(dispatch_router, prefix="/api/dispatch", tags=["Dispatch"])
app.include_router(tracking_router, prefix="/api/tracking", tags=["Tracking"])
app.include_router(billing_router, prefix="/api/billing", tags=["Billing"])
app.include_router(import_router, prefix="/api/import", tags=["Import"])
app.include_router(export_router, prefix="/api/export", tags=["Export"])
app.include_router(dashboard_router, prefix="/api/dashboard", tags=["Dashboard"])
app.include_router(audit_router, prefix="/api/audit", tags=["Audit"])
app.include_router(datastore_router, prefix="/api/datastore", tags=["Datastore"])

# ── Destination Quick Edit API ──────────────────────────────────────────────
from pydantic import BaseModel
from typing import Optional
from datetime import date, datetime
from fastapi import Depends, HTTPException
from sqlalchemy.orm import Session
from database import get_db, DB_PATH
from models import EventDestination, Dispatch, DeliveryConfirmation, DispatchItem
from routes.audit import create_audit_log
import json
import re

def normalize_region(val: Optional[str]) -> Optional[str]:
    if not val:
        return val
    s = val.strip()
    s = re.sub(r'\s*-\s*', '-', s)
    s = re.sub(r'^South\s*(\d)$', r'South-\1', s, flags=re.I)
    if s.upper() == 'BHNE':
        s = 'BNE'
    if s.upper() == 'ODISHA':
        s = 'ODC'
    if s.upper() in ('UPK', 'UPL'):
        s = 'Central'
    return s

class DestinationQuickPatch(BaseModel):
    contact_person: Optional[str] = None
    mobile: Optional[str] = None
    event_date: Optional[str] = None
    stage: Optional[str] = None  # PRINTED, IN_TRANSIT, DELIVERED, CONFIRMED
    courier_partner: Optional[str] = None
    awb_number: Optional[str] = None
    dispatch_date: Optional[str] = None
    delivery_date: Optional[str] = None
    box_count: Optional[str] = None
    influencer_cards: Optional[int] = None
    special_cards: Optional[int] = None
    team_cards: Optional[int] = None
    lanyards: Optional[int] = None

@app.patch("/api/destinations/{destination_id}")
def quick_update_destination(destination_id: int, payload: DestinationQuickPatch, db: Session = Depends(get_db)):
    dest = db.query(EventDestination).filter(EventDestination.id == destination_id).first()
    if not dest:
        raise HTTPException(status_code=404, detail="Destination not found")
    
    old_vals = {
        "contact_person": dest.contact_person,
        "mobile": dest.mobile,
        "event_date": str(dest.event_date) if dest.event_date else None
    }
    
    if payload.contact_person is not None:
        dest.contact_person = payload.contact_person.strip() or None
    if payload.mobile is not None:
        dest.mobile = payload.mobile.strip() or None
    if payload.event_date is not None:
        ed_str = payload.event_date.strip()
        if ed_str:
            try:
                dest.event_date = date.fromisoformat(ed_str)
            except ValueError:
                raise HTTPException(status_code=400, detail="Invalid date format. Use YYYY-MM-DD.")
        else:
            dest.event_date = None

    # Quantities update
    if (payload.influencer_cards is not None or payload.special_cards is not None or 
        payload.team_cards is not None or payload.lanyards is not None):
        q = dest.required_quantities or {}
        if payload.influencer_cards is not None:
            q["ITEM_PRIMARY"] = max(0, payload.influencer_cards)
            q["ITEM_PRIMARY"] = max(0, payload.influencer_cards)
        if payload.special_cards is not None:
            q["ITEM_SECONDARY"] = max(0, payload.special_cards)
            q["ITEM_SECONDARY"] = max(0, payload.special_cards)
        if payload.team_cards is not None:
            q["ITEM_VIP"] = max(0, payload.team_cards)
            q["ITEM_VIP"] = max(0, payload.team_cards)
        if payload.lanyards is not None:
            q["ITEM_ACCESSORY"] = max(0, payload.lanyards)
            q["ITEM_ACCESSORY"] = max(0, payload.lanyards)
        dest.required_quantities_json = json.dumps(q)

    # Check or find latest dispatch
    latest_dsp = db.query(Dispatch).filter(Dispatch.destination_id == dest.id).order_by(Dispatch.id.desc()).first()

    has_logistics_input = (payload.courier_partner is not None or payload.awb_number is not None or 
                           payload.dispatch_date is not None or payload.delivery_date is not None or 
                           payload.box_count is not None or payload.stage is not None)

    if has_logistics_input:
        if not latest_dsp:
            # Create a dispatch record if user provided courier info or stage beyond PRINTED
            target_stage = (payload.stage or "PRINTED").upper()
            if target_stage != "PRINTED" or payload.courier_partner or payload.awb_number:
                dsp_code = f"DSP-{dest.id}-1"
                latest_dsp = Dispatch(
                    dispatch_code=dsp_code,
                    event_id=dest.event_id,
                    destination_id=dest.id,
                    hist_organiser_name=dest.destination_name,
                    hist_contact_person=dest.contact_person,
                    hist_mobile=dest.mobile,
                    hist_address_line1=dest.address_line1,
                    hist_district=dest.district,
                    hist_city=dest.city,
                    hist_state=dest.state,
                    hist_pin=dest.pin,
                    dispatch_date=date.today(),
                    courier_partner="Trackon / Surface",
                    awb_number="AWB Pending",
                    courier_status="PENDING",
                    box_count="1 Box"
                )
                db.add(latest_dsp)
                db.flush()

        if latest_dsp:
            if payload.courier_partner is not None:
                latest_dsp.courier_partner = payload.courier_partner.strip() or "Express Courier"
            if payload.awb_number is not None:
                latest_dsp.awb_number = payload.awb_number.strip() or None
            if payload.box_count is not None:
                latest_dsp.box_count = payload.box_count.strip() or "1 Box"
            if payload.dispatch_date is not None:
                dp_str = payload.dispatch_date.strip()
                latest_dsp.dispatch_date = date.fromisoformat(dp_str) if dp_str else None
            if payload.delivery_date is not None:
                dl_str = payload.delivery_date.strip()
                latest_dsp.actual_delivery_date = date.fromisoformat(dl_str) if dl_str else None

            # Sync dispatch items if quantities were updated
            if (payload.influencer_cards is not None or payload.special_cards is not None or 
                payload.team_cards is not None or payload.lanyards is not None):
                for old_item in list(latest_dsp.items):
                    db.delete(old_item)
                db.flush()
                inf_q = payload.influencer_cards if payload.influencer_cards is not None else ((dest.required_quantities or {}).get("ITEM_PRIMARY") or (dest.required_quantities or {}).get("ITEM_PRIMARY", 0))
                spc_q = payload.special_cards if payload.special_cards is not None else ((dest.required_quantities or {}).get("ITEM_SECONDARY") or (dest.required_quantities or {}).get("ITEM_SECONDARY", 0))
                tem_q = payload.team_cards if payload.team_cards is not None else ((dest.required_quantities or {}).get("ITEM_VIP") or (dest.required_quantities or {}).get("ITEM_VIP", 0))
                lan_q = payload.lanyards if payload.lanyards is not None else ((dest.required_quantities or {}).get("ITEM_ACCESSORY") or (dest.required_quantities or {}).get("ITEM_ACCESSORY", 0))
                if inf_q > 0:
                    db.add(DispatchItem(dispatch_id=latest_dsp.id, collateral_code="ITEM_PRIMARY", quantity=inf_q))
                    db.add(DispatchItem(dispatch_id=latest_dsp.id, collateral_code="ITEM_PRIMARY", quantity=inf_q))
                if spc_q > 0:
                    db.add(DispatchItem(dispatch_id=latest_dsp.id, collateral_code="ITEM_SECONDARY", quantity=spc_q))
                    db.add(DispatchItem(dispatch_id=latest_dsp.id, collateral_code="ITEM_SECONDARY", quantity=spc_q))
                if tem_q > 0:
                    db.add(DispatchItem(dispatch_id=latest_dsp.id, collateral_code="ITEM_VIP", quantity=tem_q))
                    db.add(DispatchItem(dispatch_id=latest_dsp.id, collateral_code="ITEM_VIP", quantity=tem_q))
                if lan_q > 0:
                    db.add(DispatchItem(dispatch_id=latest_dsp.id, collateral_code="ITEM_ACCESSORY", quantity=lan_q))
                    db.add(DispatchItem(dispatch_id=latest_dsp.id, collateral_code="ITEM_ACCESSORY", quantity=lan_q))

            # Handle lifecycle stage
            if payload.stage:
                stg = payload.stage.upper()
                if stg == "CONFIRMED":
                    latest_dsp.courier_status = "DELIVERED"
                    if not latest_dsp.actual_delivery_date:
                        latest_dsp.actual_delivery_date = date.today()
                    # Delivery confirmation
                    conf = db.query(DeliveryConfirmation).filter(DeliveryConfirmation.dispatch_id == latest_dsp.id).first()
                    if not conf:
                        conf = DeliveryConfirmation(
                            dispatch_id=latest_dsp.id,
                            confirmation_status="CONFIRMED",
                            confirmed_by_name=dest.contact_person or "Station SPOC",
                            confirmation_method="PHONE_CALL",
                            confirmation_date=datetime.utcnow(),
                            notes="Confirmed on-site via Bento Station Editor"
                        )
                        db.add(conf)
                    else:
                        conf.confirmation_status = "CONFIRMED"
                        conf.confirmation_date = datetime.utcnow()
                elif stg == "DELIVERED":
                    latest_dsp.courier_status = "DELIVERED"
                    if not latest_dsp.actual_delivery_date:
                        latest_dsp.actual_delivery_date = date.today()
                    # Revert or ensure confirmation is not confirmed
                    conf = db.query(DeliveryConfirmation).filter(DeliveryConfirmation.dispatch_id == latest_dsp.id).first()
                    if conf:
                        conf.confirmation_status = "PENDING"
                elif stg == "IN_TRANSIT":
                    latest_dsp.courier_status = "IN_TRANSIT"
                    latest_dsp.actual_delivery_date = None
                    conf = db.query(DeliveryConfirmation).filter(DeliveryConfirmation.dispatch_id == latest_dsp.id).first()
                    if conf:
                        conf.confirmation_status = "PENDING"
                elif stg == "PRINTED":
                    latest_dsp.courier_status = "PENDING"
                    latest_dsp.actual_delivery_date = None
                    conf = db.query(DeliveryConfirmation).filter(DeliveryConfirmation.dispatch_id == latest_dsp.id).first()
                    if conf:
                        conf.confirmation_status = "PENDING"

            latest_dsp.updated_at = datetime.utcnow()

    dest.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(dest)
    if latest_dsp:
        db.refresh(latest_dsp)
    
    new_vals = {
        "contact_person": dest.contact_person,
        "mobile": dest.mobile,
        "event_date": str(dest.event_date) if dest.event_date else None,
        "stage": payload.stage
    }
    create_audit_log(db, "EventDestination", dest.id, "UPDATE", f"Quick updated destination {dest.destination_name}", json.dumps(old_vals), json.dumps(new_vals))
    
    return {
        "status": "success",
        "destination_id": dest.id,
        "destination_name": dest.destination_name,
        "contact_person": dest.contact_person,
        "mobile": dest.mobile,
        "event_date": str(dest.event_date) if dest.event_date else None
    }


# ── Influencer Attendee Parser & API Endpoints ──────────────────────────────
from fastapi import UploadFile, File

def parse_influencer_spreadsheet(file_path_or_bytes, filename="", target_location=None):
    """
    Robust multi-sheet parser for Pan-India and Station-level Influencer spreadsheets.
    Detects headers across rows 1-10 with word-boundary matching.
    Extracts real Name, Category (A-H), OTP, Phone, and Location.
    """
    import openpyxl, csv, io
    results = []
    is_csv = filename.lower().endswith(".csv") or (isinstance(file_path_or_bytes, str) and file_path_or_bytes.lower().endswith(".csv"))

    if is_csv:
        try:
            if isinstance(file_path_or_bytes, bytes):
                lines = file_path_or_bytes.decode("utf-8", errors="ignore").splitlines()
            else:
                with open(file_path_or_bytes, "r", encoding="utf-8", errors="ignore") as f:
                    lines = f.readlines()
            reader = list(csv.reader(lines))
            if reader:
                header_idx = None
                headers = []
                for r_idx, row in enumerate(reader[:10]):
                    r_lower = [str(c or "").strip().lower() for c in row]
                    if any(re.search(r'\b(name|inf name|guest|influencer)\b', c) for c in r_lower):
                        header_idx = r_idx
                        headers = r_lower
                        break
                if header_idx is not None:
                    name_col = next((i for i, h in enumerate(headers) if re.search(r'\b(name|inf name|full name|guest|invitee|influencer)\b', h)), None)
                    otp_col = next((i for i, h in enumerate(headers) if re.search(r'\b(otp|code|token|qr|passcode)\b', h)), None)
                    cat_col = next((i for i, h in enumerate(headers) if re.search(r'\b(category|cat|tier|grade)\b', h)), None)
                    loc_col = next((i for i, h in enumerate(headers) if re.search(r'\b(event location|location|station|venue|city|district)\b', h)), None)
                    phone_col = next((i for i, h in enumerate(headers) if re.search(r'\b(phone|mobile|contact)\b', h)), None)

                    for row in reader[header_idx + 1:]:
                        if not row: continue
                        name = str(row[name_col]).strip() if name_col is not None and len(row) > name_col and row[name_col] else ""
                        if not name or name.lower() in ["none", "null", "nan"]: continue
                        raw_cat = str(row[cat_col]).strip().upper() if cat_col is not None and len(row) > cat_col and row[cat_col] else "H"
                        cat = raw_cat if raw_cat in "ABCDEFGH" else "H"
                        otp = str(row[otp_col]).strip().replace(".0", "") if otp_col is not None and len(row) > otp_col and row[otp_col] else ""
                        if otp.lower() in ["none", "null", "nan", "#n/a"]: otp = ""
                        raw_loc = str(row[loc_col]).strip() if loc_col is not None and len(row) > loc_col and row[loc_col] else ""
                        phone = str(row[phone_col]).strip().replace(".0", "") if phone_col is not None and len(row) > phone_col and row[phone_col] else ""
                        if phone.lower() in ["none", "null", "nan"]: phone = ""

                        if target_location and loc_col is not None and raw_loc:
                            clean_rloc = re.sub(r'\b(coordinator|hub|center)\b', '', raw_loc.lower()).strip()
                            if target_location not in clean_rloc and clean_rloc not in target_location:
                                continue

                        results.append({
                            "name": name,
                            "category": cat,
                            "otp": otp,
                            "location": raw_loc or (target_location or "Station"),
                            "phone": phone
                        })
        except Exception as e:
            print(f"CSV Parse error: {e}")
    else:
        try:
            if isinstance(file_path_or_bytes, bytes):
                wb = openpyxl.load_workbook(io.BytesIO(file_path_or_bytes), data_only=True)
            else:
                wb = openpyxl.load_workbook(file_path_or_bytes, data_only=True)

            for sname in wb.sheetnames:
                sheet = wb[sname]
                header_idx = None
                headers = []
                for r_idx, row in enumerate(sheet.iter_rows(max_row=10, values_only=True), 1):
                    r_lower = [str(c or "").strip().lower() for c in row]
                    if any(re.search(r'\b(name|inf name|guest|influencer)\b', c) for c in r_lower):
                        header_idx = r_idx
                        headers = r_lower
                        break
                if not header_idx:
                    continue

                name_col = next((i for i, h in enumerate(headers) if re.search(r'\b(name|inf name|full name|guest|invitee|influencer)\b', h)), None)
                otp_col = next((i for i, h in enumerate(headers) if re.search(r'\b(otp|code|token|qr|passcode)\b', h)), None)
                cat_col = next((i for i, h in enumerate(headers) if re.search(r'\b(category|cat|tier|grade)\b', h)), None)
                loc_col = next((i for i, h in enumerate(headers) if re.search(r'\b(event location|location|station|venue|city|district)\b', h)), None)
                phone_col = next((i for i, h in enumerate(headers) if re.search(r'\b(phone|mobile|contact)\b', h)), None)

                for row in sheet.iter_rows(min_row=header_idx + 1, values_only=True):
                    name = str(row[name_col]).strip() if name_col is not None and row[name_col] is not None else ""
                    if not name or name.lower() in ["none", "null", "nan"]: continue
                    raw_cat = str(row[cat_col]).strip().upper() if cat_col is not None and row[cat_col] is not None else "H"
                    cat = raw_cat if raw_cat in "ABCDEFGH" else "H"
                    otp = str(row[otp_col]).strip().replace(".0", "") if otp_col is not None and row[otp_col] is not None else ""
                    if otp.lower() in ["none", "null", "nan", "#n/a"]: otp = ""
                    raw_loc = str(row[loc_col]).strip() if loc_col is not None and row[loc_col] is not None else ""
                    phone = str(row[phone_col]).strip().replace(".0", "") if phone_col is not None and row[phone_col] is not None else ""
                    if phone.lower() in ["none", "null", "nan"]: phone = ""

                    if target_location and loc_col is not None and raw_loc:
                        clean_rloc = re.sub(r'\b(coordinator|hub|center)\b', '', raw_loc.lower()).strip()
                        if target_location not in clean_rloc and clean_rloc not in target_location:
                            continue

                    results.append({
                        "name": name,
                        "category": cat,
                        "otp": otp,
                        "location": raw_loc or (target_location or "Station"),
                        "phone": phone
                    })
        except Exception as e:
            print(f"Excel Parse error: {e}")

    return results

@app.post("/api/destinations/{dest_id}/upload_influencers")
async def upload_destination_influencers(
    dest_id: int,
    file: UploadFile = File(...),
    db: Session = Depends(get_db)
):
    dest = db.query(EventDestination).filter(EventDestination.id == dest_id).first()
    if not dest:
        raise HTTPException(status_code=404, detail="Destination not found")

    content = await file.read()
    st_clean = re.sub(r'\b(coordinator|hub|center|station)\b', '', dest.destination_name.lower(), flags=re.I).strip()
    
    parsed = parse_influencer_spreadsheet(content, filename=file.filename, target_location=st_clean)
    if not parsed:
        # If target location didn't match, parse without location filter (assign all rows to this station)
        parsed = parse_influencer_spreadsheet(content, filename=file.filename, target_location=None)

    if not parsed:
        raise HTTPException(status_code=400, detail="No valid attendee records (Name, Category, OTP, etc.) could be recognized in the uploaded file.")

    # Save to SQLite database
    import sqlite3
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    cur.execute("DELETE FROM influencer_attendees WHERE destination_id = ? OR LOWER(TRIM(location_name)) = ?", (dest.id, st_clean))
    
    to_insert = [
        (dest.id, dest.destination_name, p["name"], p["category"], p["otp"], p["phone"], file.filename)
        for p in parsed
    ]
    cur.executemany("""
    INSERT INTO influencer_attendees (destination_id, location_name, name, category, otp, phone, source_file)
    VALUES (?, ?, ?, ?, ?, ?, ?)
    """, to_insert)
    conn.commit()
    conn.close()

    cat_counts = {c: 0 for c in "ABCDEFGH"}
    for p in parsed:
        cat_counts[p["category"]] = cat_counts.get(p["category"], 0) + 1

    return {
        "status": "success",
        "message": f"Successfully loaded {len(parsed)} real attendees from {file.filename}",
        "total_cards": len(parsed),
        "source_file": file.filename,
        "has_real_data": True,
        "categories": cat_counts,
        "attendees": [{
            "sl": idx + 1,
            "name": p["name"],
            "category": p["category"],
            "otp": p["otp"],
            "phone": p["phone"],
            "location": dest.destination_name
        } for idx, p in enumerate(parsed)]
    }

@app.get("/api/destinations/{dest_id}/category_breakdown")
def get_destination_category_breakdown(
    dest_id: int,
    file_path: Optional[str] = None,
    db: Session = Depends(get_db)
):
    dest = db.query(EventDestination).filter(EventDestination.id == dest_id).first()
    if not dest:
        raise HTTPException(status_code=404, detail="Destination not found")

    st_name = (dest.destination_name or "").strip()
    st_clean = re.sub(r'\b(coordinator|hub|center|station)\b', '', st_name.lower(), flags=re.I).strip()

    matched_attendees = []
    used_source = None

    # 1. If explicit file_path was passed
    if file_path and os.path.exists(file_path):
        parsed = parse_influencer_spreadsheet(file_path, filename=os.path.basename(file_path), target_location=st_clean)
        if not parsed:
            parsed = parse_influencer_spreadsheet(file_path, filename=os.path.basename(file_path), target_location=None)
        if parsed:
            used_source = os.path.basename(file_path)
            for idx, p in enumerate(parsed, 1):
                matched_attendees.append({
                    "sl": idx,
                    "name": p["name"],
                    "category": p["category"],
                    "otp": p["otp"],
                    "phone": p["phone"],
                    "location": st_name
                })

    # 2. Query SQLite influencer_attendees for real records
    if not matched_attendees:
        import sqlite3
        conn = sqlite3.connect(DB_PATH)
        conn.row_factory = sqlite3.Row
        cur = conn.cursor()
        rows = cur.execute("""
            SELECT name, category, otp, phone, location_name, source_file
            FROM influencer_attendees
            WHERE destination_id = ? OR LOWER(TRIM(location_name)) = ? OR LOWER(TRIM(location_name)) LIKE ?
            ORDER BY id ASC
        """, (dest.id, st_clean, f"%{st_clean}%")).fetchall()

        if rows:
            used_source = rows[0]["source_file"]
            for idx, r in enumerate(rows, 1):
                cat = (r["category"] or "H").strip().upper()
                if cat not in "ABCDEFGH": cat = "H"
                matched_attendees.append({
                    "sl": idx,
                    "name": r["name"] or f"Invitee {idx}",
                    "category": cat,
                    "otp": r["otp"] or "",
                    "phone": r["phone"] or "",
                    "location": r["location_name"] or st_name
                })
        conn.close()

    # 3. ZERO FAKE/RANDOM DATA: If no real records exist, return honest empty state!
    if not matched_attendees:
        return {
            "status": "success",
            "destination_id": dest.id,
            "destination_name": dest.destination_name,
            "state": dest.state,
            "total_cards": 0,
            "source_file": "Not Uploaded Yet",
            "has_real_data": False,
            "categories": {c: 0 for c in "ABCDEFGH"},
            "attendees": []
        }

    cat_counts = {c: 0 for c in "ABCDEFGH"}
    for a in matched_attendees:
        c = a.get("category", "H")
        cat_counts[c] = cat_counts.get(c, 0) + 1

    return {
        "status": "success",
        "destination_id": dest.id,
        "destination_name": dest.destination_name,
        "state": dest.state,
        "total_cards": len(matched_attendees),
        "source_file": used_source or "Uploaded File",
        "has_real_data": True,
        "categories": cat_counts,
        "attendees": matched_attendees
    }

# ── Serve static frontend ───────────────────────────────────────────────────
app.mount("/static", StaticFiles(directory=os.path.join(BASE_DIR, "static")), name="static")


# ── Lifecycle ────────────────────────────────────────────────────────────────
@app.on_event("startup")
async def startup():
    init_db()
    seed_collateral_items()
    # Check if database has destinations; if empty, populate with clean demo data
    try:
        from database import SessionLocal
        from models import EventDestination
        _db = SessionLocal()
        dest_count = _db.query(EventDestination).count()
        _db.close()
        if dest_count == 0:
            from seed_data import seed_demo_data
            seed_demo_data()
    except Exception as e:
        print(f"Startup notice: {e}")


# ── SPA entry point ─────────────────────────────────────────────────────────
@app.get("/")
@app.get("/index.html")
@app.get("/static/index.html")
async def root():
    return FileResponse(
        os.path.join(BASE_DIR, "static", "index.html"),
        headers={"Cache-Control": "no-cache, no-store, must-revalidate, max-age=0", "Pragma": "no-cache", "Expires": "0"}
    )


# ── Run with Uvicorn ─────────────────────────────────────────────────────────
if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="127.0.0.1", port=8000, reload=True)
